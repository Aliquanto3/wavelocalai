"""
Tests de l'import de documents de l'assistant documentaire (src/app/rag_upload.py), matrice de
la story 4 : vrai RAGEngine sur une collection Chroma sous tmp_path, embeddings déterministes,
sans reranker ni réseau.
Usage: python -m pytest tests/unit/test_rag_upload.py -v
"""

import os
from pathlib import Path

import pytest

from src.app.formatting import NBSP
from src.app.rag_upload import (
    REASON_DUPLICATE,
    REASON_FAILED,
    REASON_NO_TEXT,
    FileResult,
    escape_markdown,
    ingest_uploaded_file,
    ingest_uploaded_files,
)
from src.core.rag_engine import DEFAULT_EMBEDDING_MODEL, RAGEngine

ROOT_DIR = Path(__file__).resolve().parents[2]
RAG_PAGE = ROOT_DIR / "src" / "app" / "views" / "03_RAG_Knowledge.py"

# Assez long pour produire plusieurs extraits (découpage à 800 caractères).
MD_TEXT = "# Note de cadrage\n\n" + "\n\n".join(
    f"Paragraphe {i} : l'atelier porte sur les petits modèles de langage locaux. " * 4
    for i in range(12)
)
UNREADABLE_MARK = "CONTENU ILLISIBLE"


class FakeUpload:
    """Double de streamlit UploadedFile : un nom et des octets."""

    def __init__(self, name: str, content: str | bytes):
        self.name = name
        self._data = content.encode("utf-8") if isinstance(content, str) else content

    def getvalue(self) -> bytes:
        return self._data


@pytest.fixture
def engine(tmp_path, monkeypatch):
    """RAGEngine réel, Chroma sous tmp_path, embeddings simulés (jamais data/chroma)."""
    from langchain_core.embeddings import DeterministicFakeEmbedding

    from src.core.rag.models_factory import RAGModelsFactory

    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
    monkeypatch.setenv("ANONYMIZED_TELEMETRY", "False")
    monkeypatch.setattr(
        RAGModelsFactory,
        "get_embedding_model",
        staticmethod(lambda model_name, device="cpu": DeterministicFakeEmbedding(size=32)),
    )
    monkeypatch.setattr(
        RAGModelsFactory,
        "get_reranker_model",
        staticmethod(lambda model_name, device="cpu": None),
    )
    chroma_dir = tmp_path / "chroma"
    chroma_dir.mkdir()
    monkeypatch.setattr("src.core.rag.vector_store.CHROMA_DIR", chroma_dir)

    rag = RAGEngine()
    assert rag.vector_manager.persist_dir == str(chroma_dir)
    return rag


@pytest.fixture
def temp_paths(engine, monkeypatch):
    """Espionne ingest_file : chemins temporaires reçus, présents pendant l'appel."""
    seen = []
    original = engine.ingest_file

    def spy(file_path, original_filename):
        seen.append(file_path)
        assert os.path.exists(file_path), "le fichier temporaire doit exister pendant l'appel"
        return original(file_path, original_filename)

    monkeypatch.setattr(engine, "ingest_file", spy)
    return seen


@pytest.fixture
def unreadable_loader(monkeypatch):
    """Le loader TXT/MD lève une exception pour les fichiers qui contiennent UNREADABLE_MARK."""
    import src.core.rag.ingestion as ingestion

    real_loader = ingestion.TextLoader

    class FlakyLoader(real_loader):
        def load(self):
            if UNREADABLE_MARK in Path(self.file_path).read_text(encoding="utf-8"):
                raise UnicodeDecodeError("utf-8", b"\xff", 0, 1, "octet invalide")
            return super().load()

    monkeypatch.setattr(ingestion, "TextLoader", FlakyLoader)


def _assert_deleted(paths):
    assert paths, "ingest_file aurait dû être appelé"
    for path in paths:
        assert not os.path.exists(path), f"fichier temporaire non supprimé : {path}"


def test_single_markdown_file(engine, temp_paths):
    """Import réussi : extraits ajoutés, « N extraits ajoutés depuis 1 document »."""
    before = engine.get_stats()["count"]

    report = ingest_uploaded_files(engine, [FakeUpload("cadrage.md", MD_TEXT)])

    stats = engine.get_stats()
    assert report.added > 1
    assert stats["count"] == before + report.added
    assert stats["sources"] == ["cadrage.md"]
    assert report.failures == []
    assert report.success_message() == (
        f"{report.added}{NBSP}extraits ajoutés depuis 1{NBSP}document."
    )
    # Le fichier temporaire garde l'extension d'origine et se trouve dans le dossier temporaire.
    assert temp_paths[0].endswith(".md")
    _assert_deleted(temp_paths)


def test_single_short_file_uses_singular(engine):
    report = ingest_uploaded_files(engine, [FakeUpload("mot.txt", "Un seul extrait.")])
    assert report.added == 1
    assert report.success_message() == f"1{NBSP}extrait ajouté depuis 1{NBSP}document."


def test_several_files(engine, temp_paths):
    """Plusieurs fichiers : compte total et « 2 documents »."""
    files = [FakeUpload("cadrage.md", MD_TEXT), FakeUpload("notes.txt", "Compte rendu court.")]

    report = ingest_uploaded_files(engine, files)

    assert [r.name for r in report.succeeded] == ["cadrage.md", "notes.txt"]
    assert report.added == engine.get_stats()["count"]
    assert sorted(engine.get_stats()["sources"]) == ["cadrage.md", "notes.txt"]
    assert report.success_message() == (
        f"{report.added}{NBSP}extraits ajoutés depuis 2{NBSP}documents."
    )
    assert len(temp_paths) == 2
    _assert_deleted(temp_paths)


@pytest.mark.parametrize("content", ["", "   \n\n  "])
def test_empty_file(engine, temp_paths, content):
    """Fichier vide : aucun extrait, erreur nommant le fichier, pas de succès."""
    report = ingest_uploaded_files(engine, [FakeUpload("vide.txt", content)])

    assert engine.get_stats()["count"] == 0
    assert report.added == 0
    assert report.success_message() is None
    (failure,) = report.failures
    assert failure.name == "vide.txt"
    assert failure.reason == REASON_NO_TEXT
    assert "« vide.txt »" in failure.error_message()
    assert "aucun texte extrait" in failure.error_message()
    _assert_deleted(temp_paths)


def test_unreadable_file_does_not_block_others(engine, temp_paths, unreadable_loader):
    """Loader qui lève une exception : erreur pour ce fichier, les autres sont indexés."""
    files = [
        FakeUpload("abime.txt", UNREADABLE_MARK),
        FakeUpload("cadrage.md", MD_TEXT),
    ]

    report = ingest_uploaded_files(engine, files)

    stats = engine.get_stats()
    assert stats["sources"] == ["cadrage.md"]
    assert stats["count"] == report.added > 0
    (failure,) = report.failures
    assert failure.name == "abime.txt"
    assert failure.reason == REASON_FAILED
    assert "octet invalide" in failure.detail
    assert "« abime.txt »" in failure.error_message()
    assert report.success_message() == (
        f"{report.added}{NBSP}extraits ajoutés depuis 1{NBSP}document."
    )
    _assert_deleted(temp_paths)


def test_ingest_exception_is_reported_and_temp_file_removed(engine, temp_paths, monkeypatch):
    """Exception pendant l'indexation (ici Chroma) : raison lisible, trace en détail, fichier
    temporaire supprimé, fichiers suivants indexés."""
    store = engine.vector_manager.get_store()
    real_add = store.add_documents

    def add_documents(docs, **kwargs):
        if docs[0].metadata["source"] == "casse.md":
            raise RuntimeError("collection verrouillée")
        return real_add(docs, **kwargs)

    monkeypatch.setattr(store, "add_documents", add_documents)

    report = ingest_uploaded_files(
        engine, [FakeUpload("casse.md", MD_TEXT), FakeUpload("notes.txt", "Compte rendu.")]
    )

    (failure,) = report.failures
    assert failure.name == "casse.md"
    assert failure.reason == REASON_FAILED
    assert "collection verrouillée" in failure.detail
    assert "collection verrouillée" not in failure.error_message()
    assert engine.get_stats()["sources"] == ["notes.txt"]
    assert report.success_message() == f"1{NBSP}extrait ajouté depuis 1{NBSP}document."
    _assert_deleted(temp_paths)


def test_temp_file_removed_when_write_fails(engine, monkeypatch):
    """getvalue() qui lève : échec signalé, aucun fichier temporaire laissé."""
    import tempfile

    created = []
    real_named = tempfile.NamedTemporaryFile

    def named(*args, **kwargs):
        handle = real_named(*args, **kwargs)
        created.append(handle.name)
        return handle

    monkeypatch.setattr(tempfile, "NamedTemporaryFile", named)

    class BrokenUpload:
        name = "coupe.pdf"

        def getvalue(self):
            raise OSError("téléversement interrompu")

    result = ingest_uploaded_file(engine, BrokenUpload())

    assert not result.ok
    assert result.reason == REASON_FAILED
    assert created and not os.path.exists(created[0])


def test_all_files_fail(engine, unreadable_loader):
    """Tous en échec : aucune indexation, aucun message de succès, une erreur par fichier."""
    files = [FakeUpload("vide.txt", ""), FakeUpload("abime.md", UNREADABLE_MARK)]

    report = ingest_uploaded_files(engine, files)

    assert engine.get_stats()["count"] == 0
    assert report.success_message() is None
    assert [f.name for f in report.failures] == ["vide.txt", "abime.md"]
    assert [f.reason for f in report.failures] == [REASON_NO_TEXT, REASON_FAILED]


def test_search_returns_imported_source(engine):
    """Après import du .md de test, la recherche renvoie un extrait de ce fichier."""
    ingest_uploaded_files(engine, [FakeUpload("cadrage.md", MD_TEXT)])

    docs = engine.search("petits modèles de langage locaux", k=2)

    assert docs
    assert all(d.metadata["source"] == "cadrage.md" for d in docs)


def test_default_embedding_name_shared_by_engine_and_page(engine):
    """Même nom d'embedding par défaut pour le moteur et la page, donc même collection."""
    from src.core.rag.vector_store import VectorStoreManager

    assert engine.current_embedding_name == DEFAULT_EMBEDDING_MODEL
    page = RAG_PAGE.read_text(encoding="utf-8")
    assert "all-MiniLM" not in page, "la page doit utiliser DEFAULT_EMBEDDING_MODEL"
    assert page.count("DEFAULT_EMBEDDING_MODEL") >= 3  # import, initialisation, liste de repli
    fallback = VectorStoreManager(engine.embedding_model, DEFAULT_EMBEDDING_MODEL)
    assert fallback.collection_name == engine.vector_manager.collection_name


def test_page_has_no_fake_ingestion():
    assert "time.sleep" not in RAG_PAGE.read_text(encoding="utf-8")


def test_non_utf8_text_file(engine, temp_paths):
    """.txt non UTF-8 : échec « lecture ou indexation impossible » avec le détail technique."""
    report = ingest_uploaded_files(
        engine, [FakeUpload("latin1.txt", "Données é".encode("latin-1"))]
    )

    (failure,) = report.failures
    assert failure.reason == REASON_FAILED
    assert failure.detail
    assert engine.get_stats()["count"] == 0
    _assert_deleted(temp_paths)


def test_real_docx_file(engine, tmp_path):
    """Un vrai .docx (python-docx) est indexé : docx2txt est installé."""
    import docx

    document = docx.Document()
    document.add_heading("Note de cadrage", level=1)
    for i in range(8):
        document.add_paragraph(f"Paragraphe {i} : l'atelier porte sur les modèles locaux. " * 5)
    path = tmp_path / "cadrage.docx"
    document.save(path)

    report = ingest_uploaded_files(engine, [FakeUpload("cadrage.docx", path.read_bytes())])

    assert report.failures == [], [f.detail for f in report.failures]
    assert report.added > 0
    assert engine.get_stats()["sources"] == ["cadrage.docx"]
    docs = engine.search("modèles locaux", k=1)
    assert "modèles locaux" in docs[0].page_content


def test_temp_path_not_stored_in_metadata(engine):
    """Le chemin temporaire (supprimé, nom d'utilisateur sous Windows) n'est pas conservé."""
    ingest_uploaded_files(engine, [FakeUpload("cadrage.md", MD_TEXT)])

    metadatas = engine.vector_manager.get_store().get()["metadatas"]
    assert metadatas
    assert all("file_path" not in m for m in metadatas)
    assert all(m["source"] == "cadrage.md" for m in metadatas)


def test_data_dir_path_kept_in_metadata(tmp_path, monkeypatch):
    """Un fichier de DATA_DIR garde son chemin dans les métadonnées."""
    from src.core.rag.ingestion import IngestionPipeline

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setattr("src.core.rag.ingestion.DATA_DIR", data_dir)
    path = data_dir / "note.txt"
    path.write_text("Texte de la note.", encoding="utf-8")

    (chunk,) = IngestionPipeline().process_file(str(path), "note.txt")

    assert chunk.metadata["file_path"] == str(path.resolve())
    assert chunk.metadata["source"] == "note.txt"


def test_unsupported_extension_returns_no_chunks(tmp_path):
    """Extension non gérée : liste vide, pas d'exception."""
    from src.core.rag.ingestion import IngestionPipeline

    path = tmp_path / "script.exe"
    path.write_text("binaire", encoding="utf-8")
    assert IngestionPipeline().process_file(str(path), "script.exe") == []


def test_reimport_is_refused(engine, temp_paths):
    """Réimport d'un document déjà présent : non indexé, raison explicite, compte inchangé."""
    first = ingest_uploaded_files(engine, [FakeUpload("cadrage.md", MD_TEXT)])
    count = engine.get_stats()["count"]
    ingested = len(temp_paths)

    second = ingest_uploaded_files(
        engine, [FakeUpload("cadrage.md", MD_TEXT), FakeUpload("cadrage.md", MD_TEXT)]
    )

    assert first.added == count
    assert engine.get_stats()["count"] == count
    assert second.success_message() is None
    assert [f.reason for f in second.failures] == [REASON_DUPLICATE, REASON_DUPLICATE]
    assert "déjà présent dans la base documentaire" in second.failures[0].error_message()
    assert len(temp_paths) == ingested, "aucun fichier temporaire pour un doublon"


def test_same_name_twice_in_one_import(engine):
    """Deux fichiers du même nom dans un import : le second est refusé."""
    report = ingest_uploaded_files(
        engine, [FakeUpload("note.txt", "Première."), FakeUpload("note.txt", "Seconde.")]
    )
    assert [r.reason for r in report.results] == [None, REASON_DUPLICATE]
    assert engine.get_stats()["count"] == 1


def test_callbacks_follow_each_file(engine):
    """on_start avant chaque fichier, on_result après, avec le rapport à jour."""
    events = []

    def on_start(name):
        events.append(("start", name))

    def on_result(result, report):
        events.append(("result", result.name, result.ok, len(report.results)))

    report = ingest_uploaded_files(
        engine,
        [FakeUpload("vide.txt", ""), FakeUpload("note.txt", "Texte.")],
        on_start=on_start,
        on_result=on_result,
    )

    assert events == [
        ("start", "vide.txt"),
        ("result", "vide.txt", False, 1),
        ("start", "note.txt"),
        ("result", "note.txt", True, 2),
    ]
    assert len(report.results) == 2


def test_file_result_summary():
    assert FileResult("a.md", added=1).summary() == f"1{NBSP}extrait ajouté"
    assert FileResult("a.md", added=12).summary() == f"12{NBSP}extraits ajoutés"
    assert FileResult("a.md", reason=REASON_NO_TEXT).summary() == "aucun texte extrait"
    assert FileResult("a.md", reason=REASON_FAILED).summary() == (
        "lecture ou indexation impossible"
    )
    assert FileResult("a.md", reason=REASON_DUPLICATE).summary() == (
        "déjà présent dans la base documentaire"
    )


def test_file_name_is_escaped_in_messages():
    """Nom de fichier interpolé dans du Markdown : caractères actifs échappés."""
    name = "rapport_*final*[v2]`$x$`.md"
    assert escape_markdown(name) == r"rapport\_\*final\*\[v2\]\`\$x\$\`.md"
    message = FileResult(name, reason=REASON_NO_TEXT).error_message()
    assert f"« {escape_markdown(name)} »" in message
    assert escape_markdown("note-1.2.txt") == "note-1.2.txt"
