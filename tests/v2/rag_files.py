"""Fichiers de test de la base documentaire : un PDF avec texte, un PDF protégé, un DOCX.

Rien n'est téléchargé : le PDF est écrit à la main (une police standard, un flux de texte par
page), le PDF protégé par `pypdf.PdfWriter.encrypt` (RC4, qui n'exige pas `cryptography`), le DOCX
par `python-docx`.
"""

from __future__ import annotations

import io
from collections.abc import Sequence


def make_pdf(pages: Sequence[str]) -> bytes:
    """PDF valide, une page par texte (ASCII, sans parenthèse)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # rempli après les pages
    pages_obj = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page = add(
            f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {content} 0 R >>".encode()
        )
        kids.append(f"{page} 0 R")
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
    objects[pages_obj - 1] = (
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_blank_pdf(password: str | None = None) -> bytes:
    """PDF d'une page vide, protégé par `password` s'il est donné (RC4-128)."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password is not None:
        writer.encrypt(password, algorithm="RC4-128")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_encrypted_text_pdf(password: str = "secret") -> bytes:
    """PDF avec texte, protégé par un mot de passe (RC4-128)."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["Secret content"]))))
    writer.encrypt(password, algorithm="RC4-128")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
