# src/core/tools/document_tools.py
"""Outils de génération de documents."""

import datetime
import logging
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from langchain_core.tools import tool

from src.core.config import OUTPUT_DIR

logger = logging.getLogger(__name__)

TOOL_METADATA_DOC = {
    "name": "📝 Document Generator",
    "description": "Création de DOCX",
    "category": "output",
    "requires_config": False,
}

TOOL_METADATA_MD = {
    "name": "📋 Markdown Report",
    "description": "Rapports MD",
    "category": "output",
    "requires_config": False,
}


def _sanitize_filename(name: str) -> str:
    """Nettoie un nom de fichier."""
    name = unicodedata.normalize("NFKD", name)
    name = name.encode("ascii", "ignore").decode("ascii")
    name = "".join(c if c.isalnum() or c in "._- " else "_" for c in name)
    return name.strip()[:50] or "document"


@tool
def generate_document(title: str, content: str, author: str = "WaveLocalAI") -> str:
    """
    Génère un document Word (.docx).

    Args:
        title: Titre du document
        content: Contenu principal (supporte le Markdown basique)
        author: Auteur du document

    Returns:
        Chemin vers le document généré
    """
    try:
        doc = Document()

        # Titre
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Métadonnées
        doc.add_paragraph(f"Auteur: {author}")
        doc.add_paragraph(f"Date: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph("―" * 50)

        # Contenu
        for paragraph in content.split("\n\n"):
            if paragraph.strip():
                if paragraph.startswith("# "):
                    doc.add_heading(paragraph[2:], level=1)
                elif paragraph.startswith("## "):
                    doc.add_heading(paragraph[3:], level=2)
                elif paragraph.startswith("### "):
                    doc.add_heading(paragraph[4:], level=3)
                elif paragraph.startswith("- "):
                    for line in paragraph.split("\n"):
                        if line.startswith("- "):
                            doc.add_paragraph(line[2:], style="List Bullet")
                else:
                    doc.add_paragraph(paragraph)

        filename = f"{_sanitize_filename(title)}.docx"
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = OUTPUT_DIR / filename
        doc.save(output_path)

        return f"Document généré: {output_path}"

    except Exception as e:
        return f"Erreur génération document: {e}"


@tool
def generate_markdown_report(title: str, sections: str) -> str:
    """
    Génère un rapport au format Markdown.

    Args:
        title: Titre du rapport
        sections: Contenu des sections (format: "Section1: contenu1 | Section2: contenu2")

    Returns:
        Chemin vers le fichier Markdown généré
    """
    try:
        lines = [
            f"# {title}",
            f"*Généré le {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}*",
            "",
            "---",
            "",
        ]
        for section in sections.split("|"):
            section = section.strip()
            if ":" in section:
                section_title, section_content = section.split(":", 1)
                lines.append(f"## {section_title.strip()}")
                lines.append("")
                lines.append(section_content.strip())
                lines.append("")

        filename = f"{_sanitize_filename(title)}.md"
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = OUTPUT_DIR / filename
        output_path.write_text("\n".join(lines), encoding="utf-8")

        return f"Rapport généré: {output_path}"

    except Exception as e:
        return f"Erreur génération rapport: {e}"
