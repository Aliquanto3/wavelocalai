"""Outils de la v1 qui lisent ou écrivent un fichier du dossier de travail (story 6.2).

Copiés de `src/core/agent_tools.py` (`_analyze_csv_impl`, `_generate_*_impl`) puis adaptés : un nom
relatif au dossier de travail à la place d'un chemin libre (`Workspace`, qui refuse `..`, chemins
absolus, liens et sorties de la racine), des fichiers produits à plat dans cette racine, le nom (et non
le chemin du disque) dans le résultat. Le CSV est lu avec pandas pour rendre exactement les tableaux
de la v1 (`to_string`, `describe`) ; Excel n'est plus pris en charge (écart consigné dans la spec).
Fonctions synchrones, passées à un fil par `builtin.py` ; `pandas`, `docx` et `matplotlib` ne sont
importés que dans les fonctions.
"""

from __future__ import annotations

import datetime
import json
import re
from pathlib import Path
from typing import Any

from wavelocalai.agents.tools.registry import ToolError
from wavelocalai.agents.tools.workspace import Workspace

CHART_TYPES = ("bar", "line", "pie")
MAX_CHART_POINTS = 1000


def _stamp() -> str:
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def _reason(exc: Exception) -> str:
    """Motif d'un échec sans le chemin du disque : une `OSError` rend seulement son `strerror` (le
    chemin absolu du dossier de données n'a pas à parvenir au modèle ni à l'écran)."""
    if isinstance(exc, OSError):
        return exc.strerror or type(exc).__name__
    return str(exc)


def _generated_on() -> str:
    return datetime.datetime.now().strftime("%d/%m/%Y à %H:%M")


def analyze_csv(workspace: Workspace, filepath: str, query: str, *, max_bytes: int) -> str:
    """Analyse d'un CSV du dossier de travail (aperçu, statistiques, moyenne, somme, comptage)."""
    import pandas as pd

    path = workspace.read_path(filepath, max_bytes=max_bytes)
    if path.suffix.lower() != ".csv":
        raise ToolError("Erreur : Format non supporté. Utilisez .csv")
    try:
        df = pd.read_csv(path)
        info = f"**Analyse de {path.name}**\n\n"
        info += f"- Lignes : {len(df)}\n"
        info += f"- Colonnes : {len(df.columns)}\n"
        info += f"- Colonnes disponibles : {', '.join(str(col) for col in df.columns.tolist())}\n\n"
        query_lower = str(query).lower()
        if any(word in query_lower for word in ["aperçu", "preview", "affiche", "montre", "show"]):
            info += "**Aperçu (5 premières lignes) :**\n"
            info += df.head().to_string()
        elif any(word in query_lower for word in ["stats", "statistiques", "describe", "résumé"]):
            info += "**Statistiques descriptives :**\n"
            info += df.describe().to_string()
        elif "moyenne" in query_lower or "mean" in query_lower:
            for col in df.columns:
                if str(col).lower() in query_lower:
                    if pd.api.types.is_numeric_dtype(df[col]):
                        info += f"**Moyenne de '{col}' :** {df[col].mean():.2f}"
                    else:
                        info += f"Erreur : La colonne '{col}' n'est pas numérique"
                    break
            else:
                info += (
                    "Attention : Aucune colonne spécifique détectée. Voici les moyennes de "
                    "toutes les colonnes numériques :\n"
                )
                info += df.mean(numeric_only=True).to_string()
        elif "somme" in query_lower or "sum" in query_lower or "total" in query_lower:
            for col in df.columns:
                if str(col).lower() in query_lower:
                    if pd.api.types.is_numeric_dtype(df[col]):
                        info += f"**Somme de '{col}' :** {df[col].sum():.2f}"
                    else:
                        info += f"Erreur : La colonne '{col}' n'est pas numérique"
                    break
            else:
                info += "Attention : Aucune colonne spécifique détectée. Voici les sommes :\n"
                info += df.sum(numeric_only=True).to_string()
        elif "compte" in query_lower or "count" in query_lower or "nombre" in query_lower:
            for col in df.columns:
                if str(col).lower() in query_lower:
                    info += f"**Comptage de '{col}' :**\n{df[col].value_counts().to_string()}"
                    break
            else:
                info += f"**Nombre total de lignes :** {len(df)}"
        else:
            info += (
                "Attention : Requête non comprise. Reformulez avec : 'aperçu', "
                "'moyenne de X', 'somme de X', 'stats', etc."
            )
        return info
    except pd.errors.EmptyDataError:
        raise ToolError("Erreur : Fichier vide") from None
    except Exception as exc:  # noqa: BLE001 - l'erreur d'analyse est le résultat
        raise ToolError(f"Erreur lors de l'analyse : {_reason(exc)}") from None


def generate_document(
    workspace: Workspace, title: str, content: str, filename: str | None = None
) -> tuple[str, Path]:
    """Document Word (.docx) à partir d'un contenu Markdown simple."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    target = workspace.output_path(filename, stem="document", extension=".docx", stamp=_stamp())
    try:
        doc = Document()
        doc.add_heading(str(title), level=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Généré le {_generated_on()}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
        for raw in str(content).split("\n"):
            line = raw.strip()
            if not line:
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith(("- ", "* ")):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif re.match(r"^\d+\.\s", line):
                doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
            else:
                doc.add_paragraph(line)
        doc.save(str(target))
    except Exception as exc:  # noqa: BLE001 - l'erreur de création est le résultat
        raise ToolError(f"Erreur lors de la création du document : {_reason(exc)}") from None
    return f"Document créé : {target.name}", target


def _chart_data(data_json: Any) -> tuple[list[Any], list[Any]]:
    try:
        data = data_json if isinstance(data_json, dict) else json.loads(data_json)
    except (TypeError, json.JSONDecodeError):
        raise ToolError("Erreur : Format JSON invalide") from None
    if not isinstance(data, dict) or "labels" not in data or "values" not in data:
        raise ToolError(
            'Erreur : Format JSON invalide. Utilisez : {"labels": [...], "values": [...]}'
        )
    labels, values = data["labels"], data["values"]
    if not isinstance(labels, list) or not isinstance(values, list):
        raise ToolError(
            'Erreur : Format JSON invalide. Utilisez : {"labels": [...], "values": [...]}'
        )
    if len(labels) != len(values):
        raise ToolError("Erreur : Le nombre de labels et de valeurs doit être identique")
    if len(values) > MAX_CHART_POINTS:
        raise ToolError(f"Erreur : Trop de points (max {MAX_CHART_POINTS})")
    return labels, values


def generate_chart(
    workspace: Workspace,
    data_json: Any,
    chart_type: str = "bar",
    title: str = "Chart",
    filename: str | None = None,
) -> tuple[str, Path]:
    """Graphique PNG (barres, courbe ou camembert). L'API objet de matplotlib (`Figure`, sans
    `pyplot`) garde l'état global hors de la sortie : plusieurs fils ne se marchent pas dessus."""
    labels, values = _chart_data(data_json)
    if chart_type not in CHART_TYPES:
        raise ToolError(
            f"Erreur : Type de graphique non supporté : {chart_type}. Utilisez 'bar', 'line' ou 'pie'"
        )
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    target = workspace.output_path(filename, stem="chart", extension=".png", stamp=_stamp())
    try:
        fig = Figure(figsize=(10, 6))
        FigureCanvasAgg(fig)
        ax = fig.subplots()
        if chart_type == "bar":
            ax.bar(labels, values, color="#3498db")
            ax.set_ylabel("Valeurs")
        elif chart_type == "line":
            ax.plot(labels, values, marker="o", linewidth=2, color="#2ecc71")
            ax.set_ylabel("Valeurs")
            ax.grid(True, alpha=0.3)
        else:
            ax.pie(values, labels=labels, autopct="%1.1f%%", startangle=90)
        ax.set_title(str(title), fontsize=14, fontweight="bold")
        fig.tight_layout()
        fig.savefig(target, format="png", dpi=150, bbox_inches="tight")
    except Exception as exc:  # noqa: BLE001 - l'erreur de création est le résultat
        raise ToolError(f"Erreur lors de la création du graphique : {_reason(exc)}") from None
    return f"Graphique créé : {target.name}", target


def generate_markdown_report(
    workspace: Workspace, title: str, sections: Any, filename: str | None = None
) -> tuple[str, Path]:
    """Rapport Markdown : `sections` est un objet JSON `{titre: texte}` (ou un texte seul)."""
    if isinstance(sections, dict):
        sections_dict = sections
    else:
        text = str(sections)
        try:
            sections_dict = json.loads(text) if text.startswith("{") else {"Contenu": text}
        except json.JSONDecodeError:
            raise ToolError("Erreur : Format JSON invalide pour les sections") from None
    if not isinstance(sections_dict, dict):
        raise ToolError("Erreur : Format JSON invalide pour les sections")
    target = workspace.output_path(filename, stem="report", extension=".md", stamp=_stamp())
    report = f"# {title}\n\n*Généré le {_generated_on()}*\n\n---\n\n"
    for section_title, section_content in sections_dict.items():
        report += f"## {section_title}\n\n{section_content}\n\n"
    try:
        with target.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(report)
    except OSError as exc:
        raise ToolError(f"Erreur lors de la création du rapport : {_reason(exc)}") from None
    return f"Rapport Markdown créé : {target.name}", target
